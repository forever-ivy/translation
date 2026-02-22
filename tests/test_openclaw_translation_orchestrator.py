#!/usr/bin/env python3

import os
import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from scripts.openclaw_translation_orchestrator import (
    _agent_call,
    _available_slots,
    _strip_redundant_glossary_suffixes,
    _compact_knowledge_context,
    _compact_xlsx_prompt_payload,
    _count_xlsx_prompt_rows,
    _cap_xlsx_prompt_rows,
    _collect_translated_xlsx_keys,
    _codex_generate,
    _estimate_spreadsheet_minutes_from_candidates,
    _infer_language_pair_from_context,
    _llm_intent,
    _trim_xlsx_prompt_text,
    _validate_format_preserve_coverage,
    run,
)


def _make_docx(path: Path, text: str) -> None:
    doc = Document()
    doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _make_xlsx(path: Path, *, sheet: str, cells: dict[str, str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for addr, value in cells.items():
        ws[addr] = value
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _agent_ok(text: dict) -> dict:
    return {
        "ok": True,
        "agent_id": "mock",
        "payload": {},
        "text": json.dumps(text, ensure_ascii=False),
    }


class OpenClawTranslationOrchestratorTest(unittest.TestCase):
    @patch("scripts.openclaw_translation_orchestrator.subprocess.run")
    def test_agent_call_enforces_thinking_level(self, mocked_run):
        mocked_run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout='{"result":{"payloads":[{"text":"{}"}]}}',
            stderr="",
        )
        out = _agent_call("translator-core", "ping", timeout_seconds=90)
        self.assertTrue(out.get("ok"))
        called_cmd = mocked_run.call_args.args[0]
        self.assertIn("--thinking", called_cmd)
        self.assertIn("high", called_cmd)
        self.assertIn("--session-id", called_cmd)
        sid = called_cmd[called_cmd.index("--session-id") + 1]
        self.assertTrue(str(sid).startswith("runtime-translator-core-"))

    @patch("scripts.openclaw_translation_orchestrator.subprocess.run")
    def test_agent_call_parses_embedded_stdout_with_log_prefix(self, mocked_run):
        # Newer OpenClaw runtimes may emit a non-JSON log line prefix and return a
        # top-level payload structure (no "result" wrapper).
        stdout = (
            "[agent/embedded] google tool schema snapshot\n"
            "{\n"
            "  \"payloads\": [\n"
            "    {\n"
            "      \"text\": \"{\\\"hello\\\": \\\"world\\\"}\",\n"
            "      \"mediaUrl\": null\n"
            "    }\n"
            "  ],\n"
            "  \"meta\": {\"durationMs\": 1}\n"
            "}\n"
        )
        mocked_run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=stdout,
            stderr="",
        )
        out = _agent_call("translator-core", "ping", timeout_seconds=90)
        self.assertTrue(out.get("ok"))
        self.assertEqual(out.get("text"), '{"hello": "world"}')

    @patch("scripts.openclaw_translation_orchestrator.subprocess.run")
    def test_agent_call_marks_model_request_too_large_as_error(self, mocked_run):
        mocked_run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout='{"result":{"payloads":[{"text":"LLM request rejected: total message size 3400986 exceeds limit 2097152"}]}}',
            stderr="",
        )
        out = _agent_call("translator-core", "ping", timeout_seconds=90)
        self.assertFalse(out.get("ok"))
        self.assertTrue(str(out.get("error") or "").startswith("agent_request_too_large:"))

    @patch("scripts.openclaw_translation_orchestrator.subprocess.run")
    def test_agent_call_timeout_returns_explicit_error(self, mocked_run):
        mocked_run.side_effect = subprocess.TimeoutExpired(cmd=["openclaw"], timeout=45)
        out = _agent_call("review-core", "ping", timeout_seconds=30)
        self.assertFalse(out.get("ok"))
        self.assertEqual(out.get("error"), "agent_call_timeout:review-core")

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_revision_update_reaches_review_ready(self, mocked_call):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "Translation Task"
            review = root / "Translated -EN" / "_VERIFY" / "job_1"
            ar_dir = root / "Arabic Source"
            prev_dir = root / "Previously Translated"
            _make_docx(ar_dir / "v1 استبانة.docx", "نص عربي إصدار 1")
            _make_docx(ar_dir / "v2 استبانة.docx", "نص عربي إصدار 2 مع تعديلات")
            _make_docx(prev_dir / "V1 AI Readiness Survey.docx", "English baseline v1")

            mocked_call.side_effect = [
                _agent_ok(
                    {
                        "task_type": "REVISION_UPDATE",
                        "source_language": "ar",
                        "target_language": "en",
                        "required_inputs": ["source_old", "source_new", "target_baseline"],
                        "missing_inputs": [],
                        "confidence": 0.97,
                        "reasoning_summary": "Detected AR v1+v2 and EN baseline.",
                        "estimated_minutes": 14,
                        "complexity_score": 34,
                    }
                ),
                # Codex candidate
                _agent_ok(
                    {
                        "final_text": "Final content",
                        "final_reflow_text": "Final reflow",
                        "docx_translation_map": [{"id": "p:1", "text": "Final content"}],
                        "review_brief_points": ["Review numbering."],
                        "change_log_points": ["Applied Arabic V2 changes."],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "Initial draft done.",
                    }
                ),
                # GLM candidate
                _agent_ok(
                    {
                        "final_text": "Final content (GLM)",
                        "final_reflow_text": "Final reflow (GLM)",
                        "docx_translation_map": [{"id": "p:1", "text": "Final content (GLM)"}],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "GLM candidate done.",
                    }
                ),
                # Gemini review (Codex)
                _agent_ok(
                    {
                        "findings": [],
                        "resolved": [],
                        "unresolved": [],
                        "pass": True,
                        "terminology_rate": 0.96,
                        "structure_complete_rate": 0.96,
                        "target_language_purity": 0.98,
                        "numbering_consistency": 0.97,
                        "reasoning_summary": "Looks good.",
                    }
                ),
                # Gemini review (GLM)
                _agent_ok(
                    {
                        "findings": [],
                        "resolved": [],
                        "unresolved": [],
                        "pass": True,
                        "terminology_rate": 0.95,
                        "structure_complete_rate": 0.95,
                        "target_language_purity": 0.98,
                        "numbering_consistency": 0.98,
                        "reasoning_summary": "Pass.",
                    }
                ),
                # GLM advisory review (after rounds)
                _agent_ok(
                    {
                        "findings": [],
                        "pass": True,
                        "terminology_score": 0.95,
                        "completeness_score": 0.95,
                        "naturalness_score": 0.95,
                        "reasoning_summary": "Pass.",
                    }
                ),
            ]

            meta = {
                "job_id": "job_1",
                "root_path": str(root),
                "review_dir": str(review),
                "candidate_files": [
                    {
                        "path": str(ar_dir / "v1 استبانة.docx"),
                        "name": "v1 استبانة.docx",
                        "language": "ar",
                        "version": "v1",
                        "role": "source",
                    },
                    {
                        "path": str(ar_dir / "v2 استبانة.docx"),
                        "name": "v2 استبانة.docx",
                        "language": "ar",
                        "version": "v2",
                        "role": "source",
                    },
                    {
                        "path": str(prev_dir / "V1 AI Readiness Survey.docx"),
                        "name": "V1 AI Readiness Survey.docx",
                        "language": "en",
                        "version": "v1",
                        "role": "reference_translation",
                    },
                ],
            }

            with patch.dict(os.environ, {"OPENCLAW_GLM_ENABLED": "1"}, clear=False):
                out = run(meta)
            self.assertEqual(out["status"], "review_ready")
            self.assertTrue(out["double_pass"])
            self.assertGreaterEqual(out["iteration_count"], 1)
            self.assertLessEqual(out["iteration_count"], 3)
            self.assertEqual(out["thinking_level"], "high")
            self.assertEqual((out.get("quality_report") or {}).get("thinking_level"), "high")
            self.assertIn("final_docx", out["artifacts"])
            self.assertTrue(Path(out["artifacts"]["final_docx"]).exists())

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_empty_gemini_review_degrades_to_single_model(self, mocked_call):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "Translation Task"
            review = root / "Translated -EN" / "_VERIFY" / "job_gemini_empty"
            ar_dir = root / "Arabic Source"
            prev_dir = root / "Previously Translated"
            _make_docx(ar_dir / "v1 استبانة.docx", "نص عربي إصدار 1")
            _make_docx(ar_dir / "v2 استبانة.docx", "نص عربي إصدار 2 مع تعديلات")
            _make_docx(prev_dir / "V1 AI Readiness Survey.docx", "English baseline v1")

            mocked_call.side_effect = [
                _agent_ok(
                    {
                        "task_type": "REVISION_UPDATE",
                        "source_language": "ar",
                        "target_language": "en",
                        "required_inputs": ["source_old", "source_new", "target_baseline"],
                        "missing_inputs": [],
                        "confidence": 0.97,
                        "reasoning_summary": "Detected AR v1+v2 and EN baseline.",
                        "estimated_minutes": 14,
                        "complexity_score": 34,
                    }
                ),
                # Codex candidate
                _agent_ok(
                    {
                        "final_text": "Final content",
                        "final_reflow_text": "Final reflow",
                        "docx_translation_map": [{"id": "p:1", "text": "Final content"}],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "Initial draft done.",
                    }
                ),
                # GLM candidate
                _agent_ok(
                    {
                        "final_text": "Final content (GLM)",
                        "final_reflow_text": "Final reflow (GLM)",
                        "docx_translation_map": [{"id": "p:1", "text": "Final content (GLM)"}],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "GLM candidate done.",
                    }
                ),
                # Gemini review (Codex) - empty placeholder (treated as unavailable)
                _agent_ok(
                    {
                        "findings": [],
                        "resolved": [],
                        "unresolved": [],
                        "pass": False,
                        "terminology_rate": 0.0,
                        "structure_complete_rate": 0.0,
                        "target_language_purity": 0.0,
                        "numbering_consistency": 0.0,
                        "reasoning_summary": "",
                    }
                ),
                # Gemini review (GLM) - empty placeholder (treated as unavailable)
                _agent_ok(
                    {
                        "findings": [],
                        "resolved": [],
                        "unresolved": [],
                        "pass": False,
                        "terminology_rate": 0.0,
                        "structure_complete_rate": 0.0,
                        "target_language_purity": 0.0,
                        "numbering_consistency": 0.0,
                        "reasoning_summary": "",
                    }
                ),
                # GLM advisory review (after rounds)
                _agent_ok(
                    {
                        "findings": [],
                        "pass": True,
                        "terminology_score": 0.95,
                        "completeness_score": 0.95,
                        "naturalness_score": 0.95,
                        "reasoning_summary": "Pass.",
                    }
                ),
            ]

            meta = {
                "job_id": "job_gemini_empty",
                "root_path": str(root),
                "review_dir": str(review),
                "candidate_files": [
                    {
                        "path": str(ar_dir / "v1 استبانة.docx"),
                        "name": "v1 استبانة.docx",
                        "language": "ar",
                        "version": "v1",
                        "role": "source",
                    },
                    {
                        "path": str(ar_dir / "v2 استبانة.docx"),
                        "name": "v2 استبانة.docx",
                        "language": "ar",
                        "version": "v2",
                        "role": "source",
                    },
                    {
                        "path": str(prev_dir / "V1 AI Readiness Survey.docx"),
                        "name": "V1 AI Readiness Survey.docx",
                        "language": "en",
                        "version": "v1",
                        "role": "reference_translation",
                    },
                ],
            }

            with patch.dict(os.environ, {"OPENCLAW_GLM_ENABLED": "1"}, clear=False):
                out = run(meta)
            self.assertEqual(out["status"], "review_ready")
            self.assertTrue(out["double_pass"])
            self.assertIn("degraded_single_model", out.get("status_flags") or [])

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_spreadsheet_rounds_merge_xlsx_translation_map(self, mocked_call):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "Translation Task"
            review = root / "Translated -EN" / "_VERIFY" / "job_xlsx_merge"
            src_dir = root / "Arabic Source"

            file_a = src_dir / "FileA.xlsx"
            file_b = src_dir / "FileB.xlsx"
            _make_xlsx(file_a, sheet="Interview_FDs", cells={"B2": "مرحبا", "B3": "كيف حالك"})
            _make_xlsx(file_b, sheet="Interview_FDs", cells={"B2": "مرحبا", "B3": "كيف حالك"})

            mocked_call.side_effect = [
                _agent_ok(
                    {
                        "task_type": "SPREADSHEET_TRANSLATION",
                        "source_language": "ar",
                        "target_language": "en",
                        "required_inputs": ["source_document"],
                        "missing_inputs": [],
                        "confidence": 0.97,
                        "reasoning_summary": "Detected XLSX spreadsheet translation.",
                        "estimated_minutes": 12,
                        "complexity_score": 30,
                    }
                ),
                # Round 1: only FileA translations
                _agent_ok(
                    {
                        "final_text": "",
                        "final_reflow_text": "",
                        "docx_translation_map": [],
                        "xlsx_translation_map": [
                            {"file": "FileA.xlsx", "sheet": "Interview_FDs", "cell": "B2", "text": "Hello"},
                            {"file": "FileA.xlsx", "sheet": "Interview_FDs", "cell": "B3", "text": "How are you?"},
                        ],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "Translated FileA only (incremental).",
                    }
                ),
                # Round 2: only FileB translations, orchestrator should merge with FileA
                _agent_ok(
                    {
                        "final_text": "",
                        "final_reflow_text": "",
                        "docx_translation_map": [],
                        "xlsx_translation_map": [
                            {"file": "FileB.xlsx", "sheet": "Interview_FDs", "cell": "B2", "text": "Hello"},
                            {"file": "FileB.xlsx", "sheet": "Interview_FDs", "cell": "B3", "text": "How are you?"},
                        ],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "Translated FileB only (incremental).",
                    }
                ),
            ]

            meta = {
                "job_id": "job_xlsx_merge",
                "root_path": str(root),
                "review_dir": str(review),
                "gemini_available": False,
                "candidate_files": [
                    {"path": str(file_a), "name": "FileA.xlsx", "role": "source", "language": "ar", "version": "v1"},
                    {"path": str(file_b), "name": "FileB.xlsx", "role": "source", "language": "ar", "version": "v1"},
                ],
            }

            with patch.dict(
                os.environ,
                {
                    "OPENCLAW_GLM_ENABLED": "0",
                    "OPENCLAW_VISION_QA_IN_ROUND": "0",
                    "OPENCLAW_VISION_QA_MAX_RETRIES": "0",
                },
                clear=False,
            ):
                out = run(meta)

            self.assertEqual(out["status"], "review_ready")
            self.assertTrue(out["double_pass"])
            qr = out.get("quality_report") or {}
            preserve_by_round = qr.get("preserve_coverage_by_round") or {}
            last_round_key = str(max([int(k) for k in preserve_by_round.keys()] or [1]))
            meta2 = (preserve_by_round.get(last_round_key) or {}).get("meta") or {}
            self.assertEqual(meta2.get("xlsx_expected"), 4)
            self.assertEqual(meta2.get("xlsx_got"), 4)

            merged_selected = review / ".system" / "rounds" / f"round_{last_round_key}" / "selected_output.json"
            self.assertTrue(merged_selected.exists())
            merged = json.loads(merged_selected.read_text(encoding="utf-8"))
            self.assertEqual(len(merged.get("xlsx_translation_map") or []), 4)

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_missing_inputs_status(self, mocked_call):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "Translation Task"
            review = root / "Translated -EN" / "_VERIFY" / "job_missing"
            prev_dir = root / "Previously Translated"
            _make_docx(prev_dir / "V1 AI Readiness Survey.docx", "English baseline v1")

            mocked_call.return_value = _agent_ok(
                {
                    "task_type": "REVISION_UPDATE",
                    "source_language": "ar",
                    "target_language": "en",
                    "required_inputs": ["source_old", "source_new", "target_baseline"],
                    "missing_inputs": ["source_old", "source_new"],
                    "confidence": 0.88,
                    "reasoning_summary": "Arabic files missing.",
                    "estimated_minutes": 10,
                    "complexity_score": 18,
                }
            )

            out = run(
                {
                    "job_id": "job_missing",
                    "root_path": str(root),
                    "review_dir": str(review),
                    "candidate_files": [
                        {
                            "path": str(prev_dir / "V1 AI Readiness Survey.docx"),
                            "name": "V1 AI Readiness Survey.docx",
                            "language": "en",
                            "version": "v1",
                            "role": "reference_translation",
                        }
                    ],
                },
                plan_only=False,
            )
            self.assertEqual(out["status"], "missing_inputs")
            self.assertFalse(out["ok"])

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_bilingual_proofreading_plan_only(self, mocked_call):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "Translation Task"
            review = root / "Translated -EN" / "_VERIFY" / "job_proof"
            _make_docx(root / "english_original.docx", "English original text")
            _make_docx(root / "french_translation.docx", "Texte traduit en français")

            mocked_call.return_value = _agent_ok(
                {
                    "task_type": "BILINGUAL_PROOFREADING",
                    "source_language": "en",
                    "target_language": "fr",
                    "required_inputs": ["source_document", "target_document"],
                    "missing_inputs": [],
                    "confidence": 0.92,
                    "reasoning_summary": "Proofreading EN→FR pair.",
                    "estimated_minutes": 8,
                    "complexity_score": 20,
                }
            )

            out = run(
                {
                    "job_id": "job_proof",
                    "root_path": str(root),
                    "review_dir": str(review),
                    "message_text": "could you kindly review the attached paragraphs and proofread the French version",
                    "candidate_files": [
                        {
                            "path": str(root / "english_original.docx"),
                            "name": "english_original.docx",
                            "language": "en",
                            "version": "unknown",
                            "role": "general",
                        },
                        {
                            "path": str(root / "french_translation.docx"),
                            "name": "french_translation.docx",
                            "language": "fr",
                            "version": "unknown",
                            "role": "general",
                        },
                    ],
                },
                plan_only=True,
            )
            self.assertEqual(out["status"], "planned")
            self.assertEqual(out["intent"]["task_type"], "BILINGUAL_PROOFREADING")
            self.assertGreater(out["intent"]["confidence"], 0.0)


class IntentFallbackLanguageInferenceTest(unittest.TestCase):
    def test_infer_language_pair_dedupes_candidate_languages(self):
        candidates = [
            {"language": "ar"},
            {"language": "ar"},
            {"language": "ar"},
        ]
        src, tgt = _infer_language_pair_from_context("", candidates)
        self.assertEqual(src, "ar")
        self.assertEqual(tgt, "en")

    def test_infer_language_pair_accepts_english_typos(self):
        candidates = [{"language": "ar"}]
        src, tgt = _infer_language_pair_from_context("translate arabic to englsih", candidates)
        self.assertEqual(src, "ar")
        self.assertEqual(tgt, "en")


class SpreadsheetEtaHeuristicsTest(unittest.TestCase):
    def test_estimate_spreadsheet_minutes_scales_with_size(self):
        small = [
            {
                "path": "/tmp/a.xlsx",
                "structure": {"blocks": [{"text": "A1"}, {"text": "A2"}]},
            }
        ]
        large = [
            {
                "path": "/tmp/b.xlsx",
                "structure": {
                    "blocks": [{"text": ("x" * 120)} for _ in range(220)],
                },
            }
        ]
        eta_small = _estimate_spreadsheet_minutes_from_candidates(small)
        eta_large = _estimate_spreadsheet_minutes_from_candidates(large)
        self.assertGreaterEqual(eta_small, 8)
        self.assertGreater(eta_large, eta_small)

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_llm_intent_spreadsheet_eta_uses_dynamic_floor(self, mocked_agent_call):
        mocked_agent_call.return_value = _agent_ok(
            {
                "task_type": "SPREADSHEET_TRANSLATION",
                "task_label": "Translate Arabic Excel file to English",
                "source_language": "ar",
                "target_language": "en",
                "required_inputs": ["source_document"],
                "missing_inputs": [],
                "confidence": 0.9,
                "reasoning_summary": "Spreadsheet translation",
                "estimated_minutes": 12,
                "complexity_score": 20.0,
            }
        )
        candidates = [
            {
                "path": "/tmp/large.xlsx",
                "name": "large.xlsx",
                "language": "ar",
                "structure": {
                    "blocks": [{"text": ("x" * 120)} for _ in range(220)],
                },
            }
        ]
        result = _llm_intent({"subject": "translate", "message_text": "translate arabic to english"}, candidates)
        self.assertTrue(result.get("ok"))
        self.assertEqual((result.get("intent") or {}).get("task_type"), "SPREADSHEET_TRANSLATION")
        self.assertGreaterEqual(int(result.get("estimated_minutes") or 0), 30)


class GlossarySuffixCleanupTest(unittest.TestCase):
    def test_strips_redundant_tail_glossary_label_in_xlsx_cell(self):
        context = {
            "glossary_enforcer": {
                "enabled": True,
                "terms": [{"ar": "الذكاء الاصطناعي", "en": "Artificial Intelligence (AI)"}],
            },
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "survey.xlsx",
                        "cell_units": [
                            {
                                "sheet": "Sheet1",
                                "cell": "A1",
                                "text": "ما نوع التواصل الذي تلقيته بخصوص الذكاء الاصطناعي؟",
                            }
                        ],
                    }
                ]
            },
        }
        draft = {
            "xlsx_translation_map": [
                {
                    "file": "survey.xlsx",
                    "sheet": "Sheet1",
                    "cell": "A1",
                    "text": "What communication did you receive regarding AI policies? Artificial Intelligence (AI)",
                }
            ]
        }

        meta = _strip_redundant_glossary_suffixes(context, draft)
        self.assertTrue(bool(meta.get("enabled")))
        self.assertEqual(int(meta.get("cleaned_xlsx_cells") or 0), 1)
        self.assertEqual(
            draft["xlsx_translation_map"][0]["text"],
            "What communication did you receive regarding AI policies?",
        )

    def test_keeps_non_duplicated_term(self):
        context = {
            "glossary_enforcer": {
                "enabled": True,
                "terms": [{"ar": "الذكاء الاصطناعي", "en": "Artificial Intelligence (AI)"}],
            },
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "survey.xlsx",
                        "cell_units": [{"sheet": "Sheet1", "cell": "A1", "text": "الذكاء الاصطناعي"}],
                    }
                ]
            },
        }
        draft = {
            "xlsx_translation_map": [
                {
                    "file": "survey.xlsx",
                    "sheet": "Sheet1",
                    "cell": "A1",
                    "text": "Artificial Intelligence (AI)",
                }
            ]
        }

        meta = _strip_redundant_glossary_suffixes(context, draft)
        self.assertTrue(bool(meta.get("enabled")))
        self.assertEqual(int(meta.get("cleaned_xlsx_cells") or 0), 0)
        self.assertEqual(draft["xlsx_translation_map"][0]["text"], "Artificial Intelligence (AI)")


class PromptCompactionHelpersTest(unittest.TestCase):
    def test_compact_knowledge_context_truncates_payload(self):
        hits = [
            {
                "id": "h1",
                "title": "Doc",
                "text": "A" * 5000,
                "score": 0.9,
            }
        ] * 10
        compact = _compact_knowledge_context(hits)
        self.assertGreater(len(compact), 0)
        self.assertLessEqual(len(compact), 6)
        snippet = str((compact[0] or {}).get("snippet") or "")
        self.assertTrue(len(snippet) <= 1201)

    def test_trim_xlsx_prompt_text(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "cell_units": [
                            {"file": "a.xlsx", "sheet": "S1", "cell": "A1", "text": "B" * 300},
                            {"file": "a.xlsx", "sheet": "S1", "cell": "A2", "text": "short"},
                        ],
                    }
                ]
            }
        }
        trimmed = _trim_xlsx_prompt_text(context, max_chars_per_cell=80)
        self.assertEqual(trimmed, 1)
        units = context["format_preserve"]["xlsx_sources"][0]["cell_units"]
        self.assertTrue(len(units[0]["text"]) <= 81)
        self.assertEqual(units[1]["text"], "short")

    def test_trim_xlsx_prompt_text_rows_mode(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "rows": [
                            ["S1", "A1", "B" * 260],
                            ["S1", "A2", "short"],
                        ],
                    }
                ]
            }
        }
        trimmed = _trim_xlsx_prompt_text(context, max_chars_per_cell=80)
        self.assertEqual(trimmed, 1)
        rows = context["format_preserve"]["xlsx_sources"][0]["rows"]
        self.assertTrue(len(rows[0][2]) <= 81)
        self.assertEqual(rows[1][2], "short")

    def test_compact_xlsx_prompt_payload_filters_translated_keys(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "path": "/tmp/a.xlsx",
                        "meta": {"cell_count": 2},
                        "cell_units": [
                            {"file": "a.xlsx", "sheet": "S1", "cell": "A1", "text": "first"},
                            {"file": "a.xlsx", "sheet": "S1", "cell": "A2", "text": "second"},
                        ],
                    }
                ]
            }
        }
        previous = {
            "xlsx_translation_map": [
                {"file": "a.xlsx", "sheet": "S1", "cell": "A1", "text": "FIRST"},
            ]
        }
        translated = _collect_translated_xlsx_keys(previous)
        self.assertIn(("a.xlsx", "S1", "A1"), translated)

        stats = _compact_xlsx_prompt_payload(context, previous_payload=previous)
        self.assertTrue(stats["changed"])
        self.assertEqual(stats["total_rows"], 2)
        self.assertEqual(stats["kept_rows"], 1)
        self.assertEqual(stats["skipped_existing"], 1)
        rows = context["format_preserve"]["xlsx_sources"][0]["rows"]
        self.assertEqual(rows, [["S1", "A2", "second"]])

    def test_cap_xlsx_prompt_rows(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {"file": "a.xlsx", "rows": [["S1", "A1", "x"], ["S1", "A2", "y"]]},
                    {"file": "b.xlsx", "rows": [["S2", "B1", "z"]]},
                ]
            }
        }
        self.assertEqual(_count_xlsx_prompt_rows(context), 3)
        kept = _cap_xlsx_prompt_rows(context, max_rows=2)
        self.assertEqual(kept, 2)
        self.assertEqual(_count_xlsx_prompt_rows(context), 2)

    def test_cap_xlsx_prompt_rows_cell_units(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "cell_units": [
                            {"sheet": "S1", "cell": "A1", "text": "x"},
                            {"sheet": "S1", "cell": "A2", "text": "y"},
                        ],
                    },
                    {"file": "b.xlsx", "cell_units": [{"sheet": "S2", "cell": "B1", "text": "z"}]},
                ]
            }
        }
        self.assertEqual(_count_xlsx_prompt_rows(context), 3)
        kept = _cap_xlsx_prompt_rows(context, max_rows=2)
        self.assertEqual(kept, 2)
        self.assertEqual(_count_xlsx_prompt_rows(context), 2)
        self.assertEqual(len(context["format_preserve"]["xlsx_sources"][0]["cell_units"]), 2)
        self.assertEqual(context["format_preserve"]["xlsx_sources"][1]["cell_units"], [])

    def test_validate_format_preserve_coverage_classifies_translation_vs_source_truncation(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "cell_units": [
                            {"file": "a.xlsx", "sheet": "S1", "cell": "D18", "text": "complete source sentence."},
                            {"file": "a.xlsx", "sheet": "S1", "cell": "D19", "text": "incomplete source sentence"},
                        ],
                    }
                ]
            }
        }
        draft = {
            "xlsx_translation_map": [
                {"file": "a.xlsx", "sheet": "S1", "cell": "D18", "text": "truncated output tail " * 8},
                {"file": "a.xlsx", "sheet": "S1", "cell": "D19", "text": "truncated output tail " * 8},
            ]
        }
        findings, meta = _validate_format_preserve_coverage(context, draft)
        self.assertIn("xlsx_translation_truncated:cells=1", findings)
        self.assertEqual(meta.get("xlsx_source_truncated_count"), 1)

    def test_validate_format_preserve_coverage_source_truncated_marker_not_hard_fail(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "cell_units": [
                            {"file": "a.xlsx", "sheet": "S1", "cell": "D19", "text": "incomplete source sentence و"},
                        ],
                    }
                ]
            }
        }
        draft = {
            "xlsx_translation_map": [
                {
                    "file": "a.xlsx",
                    "sheet": "S1",
                    "cell": "D19",
                    "text": ("incomplete output " * 12) + "[SOURCE TRUNCATED]",
                },
            ]
        }
        findings, meta = _validate_format_preserve_coverage(context, draft)
        self.assertEqual(findings, [])
        self.assertEqual(meta.get("xlsx_source_truncated_count"), 1)
        self.assertEqual(meta.get("xlsx_source_truncated_marker_missing_count"), 0)

    def test_validate_format_preserve_coverage_marker_is_hard_truncation_when_source_not_cutoff_like(self):
        context = {
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "a.xlsx",
                        "cell_units": [
                            {"file": "a.xlsx", "sheet": "S1", "cell": "D18", "text": "complete source sentence."},
                        ],
                    }
                ]
            }
        }
        draft = {
            "xlsx_translation_map": [
                {
                    "file": "a.xlsx",
                    "sheet": "S1",
                    "cell": "D18",
                    "text": ("truncated output " * 12) + "[SOURCE TRUNCATED]",
                },
            ]
        }
        findings, meta = _validate_format_preserve_coverage(context, draft)
        self.assertIn("xlsx_translation_truncated:cells=1", findings)
        self.assertEqual(meta.get("xlsx_source_truncated_count"), None)


class CodexGenerateFallbackTest(unittest.TestCase):
    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_spreadsheet_batches_and_merges(self, mocked_agent_call):
        units = []
        for idx in range(1, 9):
            units.append({"file": "fd.xlsx", "sheet": "S1", "cell": f"A{idx}", "text": f"نص {idx}"})

        batch1_map = [
            {"file": "fd.xlsx", "sheet": "S1", "cell": f"A{idx}", "text": f"text {idx}"}
            for idx in range(1, 7)
        ]
        batch2_map = [
            {"file": "fd.xlsx", "sheet": "S1", "cell": f"A{idx}", "text": f"text {idx}"}
            for idx in range(7, 9)
        ]

        mocked_agent_call.side_effect = [
            _agent_ok(
                {
                    "final_text": "",
                    "final_reflow_text": "",
                    "docx_translation_map": [],
                    "xlsx_translation_map": batch1_map,
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "batch1",
                }
            ),
            _agent_ok(
                {
                    "final_text": "",
                    "final_reflow_text": "",
                    "docx_translation_map": [],
                    "xlsx_translation_map": batch2_map,
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "batch2",
                }
            ),
        ]

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate",
            "candidate_files": [],
            "format_preserve": {
                "xlsx_sources": [
                    {
                        "file": "fd.xlsx",
                        "cell_units": units,
                    }
                ]
            },
        }
        with patch.dict(
            os.environ,
            {
                "OPENCLAW_XLSX_BATCH_MAX_CELLS": "6",
                "OPENCLAW_XLSX_BATCH_MAX_SOURCE_CHARS": "8000",
                "OPENCLAW_XLSX_BATCH_RETRY": "1",
                "OPENCLAW_GLM_DIRECT_FALLBACK_ENABLED": "0",
                "OPENCLAW_KIMI_CODING_DIRECT_FALLBACK_ENABLED": "0",
            },
            clear=False,
        ):
            out = _codex_generate(context, None, [], 1)

        self.assertTrue(out.get("ok"))
        data = out.get("data") or {}
        self.assertEqual(len(data.get("xlsx_translation_map") or []), 8)
        self.assertEqual(mocked_agent_call.call_count, 2)

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_spreadsheet_batch_retries_missing_cells(self, mocked_agent_call):
        units = [{"file": "fd.xlsx", "sheet": "S1", "cell": f"A{idx}", "text": f"نص {idx}"} for idx in range(1, 5)]
        missing_first = [
            {"file": "fd.xlsx", "sheet": "S1", "cell": "A1", "text": "text 1"},
            {"file": "fd.xlsx", "sheet": "S1", "cell": "A2", "text": "text 2"},
            {"file": "fd.xlsx", "sheet": "S1", "cell": "A3", "text": "text 3"},
        ]
        full_retry = missing_first + [{"file": "fd.xlsx", "sheet": "S1", "cell": "A4", "text": "text 4"}]

        mocked_agent_call.side_effect = [
            _agent_ok(
                {
                    "final_text": "",
                    "final_reflow_text": "",
                    "docx_translation_map": [],
                    "xlsx_translation_map": missing_first,
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "attempt1",
                }
            ),
            _agent_ok(
                {
                    "final_text": "",
                    "final_reflow_text": "",
                    "docx_translation_map": [],
                    "xlsx_translation_map": full_retry,
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "attempt2",
                }
            ),
        ]

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate",
            "candidate_files": [],
            "format_preserve": {
                "xlsx_sources": [{"file": "fd.xlsx", "cell_units": units}]
            },
        }
        with patch.dict(
            os.environ,
            {
                "OPENCLAW_XLSX_BATCH_MAX_CELLS": "6",
                "OPENCLAW_XLSX_BATCH_RETRY": "1",
                "OPENCLAW_GLM_DIRECT_FALLBACK_ENABLED": "0",
                "OPENCLAW_KIMI_CODING_DIRECT_FALLBACK_ENABLED": "0",
            },
            clear=False,
        ):
            out = _codex_generate(context, None, [], 1)

        self.assertTrue(out.get("ok"))
        self.assertEqual(len((out.get("data") or {}).get("xlsx_translation_map") or []), 4)
        self.assertEqual(mocked_agent_call.call_count, 2)

    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_spreadsheet_batch_splits_on_request_too_large(self, mocked_agent_call):
        units = [
            {"file": "fd.xlsx", "sheet": "S1", "cell": "A1", "text": "نص 1"},
            {"file": "fd.xlsx", "sheet": "S1", "cell": "A2", "text": "نص 2"},
        ]

        mocked_agent_call.side_effect = [
            {
                "ok": False,
                "error": "agent_request_too_large:translator-core",
                "detail": "LLM request rejected: total message size 3000000 exceeds limit 2097152",
                "raw_text": "LLM request rejected: total message size 3000000 exceeds limit 2097152",
            },
            _agent_ok(
                {
                    "final_text": "",
                    "final_reflow_text": "",
                    "docx_translation_map": [],
                    "xlsx_translation_map": [{"file": "fd.xlsx", "sheet": "S1", "cell": "A1", "text": "text 1"}],
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "split a",
                }
            ),
            _agent_ok(
                {
                    "final_text": "",
                    "final_reflow_text": "",
                    "docx_translation_map": [],
                    "xlsx_translation_map": [{"file": "fd.xlsx", "sheet": "S1", "cell": "A2", "text": "text 2"}],
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "split b",
                }
            ),
        ]

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate",
            "candidate_files": [],
            "format_preserve": {
                "xlsx_sources": [{"file": "fd.xlsx", "cell_units": units}]
            },
        }
        with patch("scripts.openclaw_translation_orchestrator.CODEX_FALLBACK_AGENT", ""), patch.dict(
            os.environ,
            {
                "OPENCLAW_XLSX_BATCH_MAX_CELLS": "10",
                "OPENCLAW_XLSX_BATCH_RETRY": "0",
                "OPENCLAW_GLM_DIRECT_FALLBACK_ENABLED": "0",
                "OPENCLAW_KIMI_CODING_DIRECT_FALLBACK_ENABLED": "0",
            },
            clear=False,
        ):
            out = _codex_generate(context, None, [], 1)

        self.assertTrue(out.get("ok"))
        self.assertEqual(len((out.get("data") or {}).get("xlsx_translation_map") or []), 2)
        self.assertEqual(mocked_agent_call.call_count, 3)

    @patch("scripts.openclaw_translation_orchestrator._kimi_coding_direct_api_call")
    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_uses_fallback_agent_on_request_too_large(self, mocked_agent_call, mocked_kimi_call):
        mocked_agent_call.side_effect = [
            {
                "ok": False,
                "error": "agent_request_too_large:translator-core",
                "detail": "LLM request rejected: total message size 3400986 exceeds limit 2097152",
                "raw_text": "LLM request rejected: total message size 3400986 exceeds limit 2097152",
            },
            {
                "ok": True,
                "agent_id": "qa-gate",
                "payload": {},
                "text": json.dumps(
                    {
                        "final_text": "English output via fallback agent",
                        "final_reflow_text": "English output via fallback agent",
                        "docx_translation_map": [],
                        "xlsx_translation_map": [],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "ok",
                    }
                ),
                "meta": {"provider": "moonshot", "model": "moonshot/kimi-k2.5"},
            },
        ]
        mocked_kimi_call.return_value = {"ok": False, "error": "should_not_be_called"}

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate arabic to english",
            "candidate_files": [],
        }
        out = _codex_generate(context, None, [], 1)
        self.assertTrue(out.get("ok"))
        self.assertEqual((out.get("call_meta") or {}).get("provider"), "moonshot")
        self.assertEqual((out.get("call_meta") or {}).get("model"), "moonshot/kimi-k2.5")
        mocked_kimi_call.assert_not_called()

    @patch("scripts.openclaw_translation_orchestrator._kimi_coding_direct_api_call")
    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_uses_fallback_agent_before_direct_api(self, mocked_agent_call, mocked_kimi_call):
        mocked_agent_call.side_effect = [
            {
                "ok": True,
                "agent_id": "translator-core",
                "payload": {},
                "text": (
                    "Cloud Code Assist API error (400): Invalid JSON payload received. "
                    "Unknown name \"patternProperties\" at "
                    "'request.tools[0].function_declarations[3].parameters.properties[2].value': Cannot find field."
                ),
                "meta": {"provider": "google-antigravity", "model": "claude-opus-4-6-thinking"},
            },
            {
                "ok": True,
                "agent_id": "qa-gate",
                "payload": {},
                "text": json.dumps(
                    {
                        "final_text": "English output via fallback agent",
                        "final_reflow_text": "English output via fallback agent",
                        "docx_translation_map": [],
                        "xlsx_translation_map": [],
                        "review_brief_points": [],
                        "change_log_points": [],
                        "resolved": [],
                        "unresolved": [],
                        "codex_pass": True,
                        "reasoning_summary": "ok",
                    }
                ),
                "meta": {"provider": "openai-codex", "model": "gpt-5.2"},
            },
        ]
        mocked_kimi_call.return_value = {"ok": False, "error": "should_not_be_called"}

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate arabic to english",
            "candidate_files": [],
        }
        out = _codex_generate(context, None, [], 1)
        self.assertTrue(out.get("ok"))
        self.assertEqual((out.get("call_meta") or {}).get("provider"), "openai-codex")
        self.assertEqual((out.get("call_meta") or {}).get("model"), "gpt-5.2")
        mocked_kimi_call.assert_not_called()

    @patch("scripts.openclaw_translation_orchestrator._kimi_coding_direct_api_call")
    @patch("scripts.openclaw_translation_orchestrator._glm_direct_api_call")
    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_falls_back_to_kimi_direct_api_on_schema_error(
        self,
        mocked_agent_call,
        mocked_glm_call,
        mocked_kimi_call,
    ):
        mocked_agent_call.return_value = {
            "ok": True,
            "agent_id": "translator-core",
            "payload": {},
            "text": (
                "Cloud Code Assist API error (400): Invalid JSON payload received. "
                "Unknown name \"patternProperties\" at "
                "'request.tools[0].function_declarations[3].parameters.properties[2].value': Cannot find field."
            ),
            "meta": {"provider": "google-antigravity", "model": "claude-opus-4-6-thinking"},
        }
        mocked_glm_call.return_value = {"ok": False, "error": "glm_direct_failed"}
        mocked_kimi_call.return_value = {
            "ok": True,
            "text": json.dumps(
                {
                    "final_text": "English output",
                    "final_reflow_text": "English output",
                    "docx_translation_map": [],
                    "xlsx_translation_map": [],
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "ok",
                }
            ),
            "source": "direct_api_kimi_coding",
            "provider": "kimi-coding",
            "model": "kimi-coding/k2p5",
        }

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate arabic to english",
            "candidate_files": [],
        }
        out = _codex_generate(context, None, [], 1)
        self.assertTrue(out.get("ok"))
        self.assertEqual((out.get("call_meta") or {}).get("provider"), "kimi-coding")
        self.assertEqual((out.get("call_meta") or {}).get("model"), "kimi-coding/k2p5")
        mocked_glm_call.assert_called_once()
        mocked_kimi_call.assert_called_once()

    @patch("scripts.openclaw_translation_orchestrator._kimi_coding_direct_api_call")
    @patch("scripts.openclaw_translation_orchestrator._glm_direct_api_call")
    @patch("scripts.openclaw_translation_orchestrator._agent_call")
    def test_codex_generate_falls_back_to_kimi_direct_api_on_json_parse_error(
        self,
        mocked_agent_call,
        mocked_glm_call,
        mocked_kimi_call,
    ):
        mocked_agent_call.return_value = {
            "ok": True,
            "agent_id": "translator-core",
            "payload": {},
            "text": "not a json payload",
            "meta": {"provider": "kimi-coding", "model": "kimi-coding/k2p5"},
        }
        mocked_glm_call.return_value = {"ok": False, "error": "glm_direct_failed"}
        mocked_kimi_call.return_value = {
            "ok": True,
            "text": json.dumps(
                {
                    "final_text": "English output via direct kimi",
                    "final_reflow_text": "English output via direct kimi",
                    "docx_translation_map": [],
                    "xlsx_translation_map": [],
                    "review_brief_points": [],
                    "change_log_points": [],
                    "resolved": [],
                    "unresolved": [],
                    "codex_pass": True,
                    "reasoning_summary": "ok",
                }
            ),
            "source": "direct_api_kimi_coding",
            "provider": "kimi-coding",
            "model": "kimi-coding/k2p5",
        }

        context = {
            "task_intent": {"task_type": "SPREADSHEET_TRANSLATION"},
            "subject": "Translate",
            "message_text": "translate arabic to english",
            "candidate_files": [],
        }
        out = _codex_generate(context, None, [], 1)
        self.assertTrue(out.get("ok"))
        self.assertEqual((out.get("call_meta") or {}).get("provider"), "kimi-coding")
        self.assertEqual((out.get("call_meta") or {}).get("model"), "kimi-coding/k2p5")
        mocked_glm_call.assert_called_once()
        mocked_kimi_call.assert_called_once()


class AvailableSlotsTest(unittest.TestCase):
    def test_french_english_pair(self):
        candidates = [
            {"language": "en", "version": "unknown", "role": "general"},
            {"language": "fr", "version": "unknown", "role": "general"},
        ]
        slots = _available_slots(candidates, source_language="en", target_language="fr")
        self.assertTrue(slots["source_document"])
        self.assertTrue(slots["target_document"])

    def test_single_arabic_file(self):
        candidates = [
            {"language": "ar", "version": "v1", "role": "source"},
        ]
        slots = _available_slots(candidates, source_language="ar", target_language="en")
        self.assertTrue(slots["source_document"])
        self.assertFalse(slots["target_document"])

    def test_empty_candidates(self):
        slots = _available_slots([], source_language="en", target_language="fr")
        self.assertFalse(slots["source_document"])
        self.assertFalse(slots["target_document"])


if __name__ == "__main__":
    unittest.main()
